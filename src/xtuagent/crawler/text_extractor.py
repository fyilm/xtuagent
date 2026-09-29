import os
import re
import logging
from pathlib import Path
from typing import Optional

import pdfplumber

logger = logging.getLogger(__name__)

try:  # python-docx 是可选依赖：缺失时只影响 .docx 解析，其余功能不受影响
    import docx as _docx
except ImportError:  # pragma: no cover - 取决于安装环境
    _docx = None


class TextExtractor:
    def __init__(self, output_dir: str = "data/texts"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _clean_text(text: str) -> str:
        text = re.sub(r"\x00", "", text)
        text = re.sub(r"\r\n", "\n", text)
        text = re.sub(r"\r", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r" +", " ", text)
        text = re.sub(r"^\s*\d+\s*$", "", text, flags=re.MULTILINE)
        return text.strip()

    def extract_pdf(self, filepath: str) -> Optional[str]:
        path = Path(filepath)
        try:
            with pdfplumber.open(path) as pdf:
                pages = []
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        pages.append(self._clean_text(text))
                return "\n\n".join(pages)
        except Exception as e:
            logger.warning("pdf extract failed: %s — %s", path, e)
            return None

    @staticmethod
    def extract_docx(filepath: str) -> Optional[str]:
        """解析 .docx（含表格单元格）。

        培养方案、学生手册这类文档很多是 Word，且关键信息（学分、课程、
        比例）经常放在表格里，所以表格必须一并抽出来。
        """
        if _docx is None:
            logger.warning("未安装 python-docx，无法解析 .docx：%s", filepath)
            return None
        try:
            document = _docx.Document(filepath)
        except Exception as e:  # noqa: BLE001
            logger.warning("docx extract failed: %s — %s", filepath, e)
            return None

        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        return TextExtractor._clean_text(text) if text else None

    def extract_and_save(self, filepath: str) -> Optional[str]:
        path = Path(filepath)
        suffix = path.suffix.lower()
        out_name = path.stem + ".txt"
        out_path = self.output_dir / out_name

        if out_path.exists():
            return str(out_path)

        if suffix == ".pdf":
            text = self.extract_pdf(str(path))
        elif suffix in (".md", ".txt"):
            try:
                text = self._clean_text(path.read_text(encoding="utf-8"))
            except (UnicodeDecodeError, Exception):
                try:
                    text = self._clean_text(path.read_text(encoding="gbk"))
                except Exception:
                    return None
        elif suffix == ".docx":
            text = self.extract_docx(str(path))
        elif suffix == ".doc":
            # 老版二进制 .doc 无纯 Python 解析器，需先另存为 .docx/.pdf
            logger.warning("legacy .doc not supported, please convert to .docx/.pdf: %s", path)
            return None
        else:
            return None

        if not text or len(text) < 50:
            return None

        out_path.write_text(text, encoding="utf-8")
        logger.info("extracted: %s -> %s", path, out_path)
        return str(out_path)

    def extract_batch(self, filepaths: list[str]) -> list[str]:
        results = []
        for fp in filepaths:
            result = self.extract_and_save(fp)
            if result:
                results.append(result)
        return results
